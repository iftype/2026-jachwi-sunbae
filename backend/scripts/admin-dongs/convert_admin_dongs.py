#!/usr/bin/env python3
"""통계청 SGIS 행정동 경계 Shapefile을 행정동 적재용 GeoJSON으로 바꾼다.

해당 스크립트는 개발자 PC에서 경계 데이터를 갱신할 때 한 번 실행하는 도구이다.
좌표계 변환과 단순화 기준이 곧 적재 데이터의 정책이므로 저장소에 남겨 둔다.
결과 파일은 약 30MB라 저장소에 커밋하지 않는다.

입력:
    - SGIS 행정구역 경계(행정동) Shapefile 압축 파일. 안에 BND_ADM_DONG_PG.* 가 있어야 한다.
    - SGIS 센서스 공간정보 지역 코드 엑셀. 시도·시군구·읍면동 코드와 이름이 3행부터 있어야 한다.

출력 GeoJSON의 Feature 속성(애플리케이션의 AdminDongGeoJsonReader가 읽는 형식):
    code(통계청 행정동 코드 8자리), name, sidoName, sigunguName, fullName, baseDate(yyyyMMdd)

필요한 패키지:
    pip install pyshp shapely pyproj openpyxl

사용법:
    python3 convert_admin_dongs.py <경계 zip> <지역 코드 xlsx> <출력 GeoJSON> [--sheet "2025년 6월"]

적재:
    java -jar app.jar --spring.main.web-application-type=none \\
        --transit.admin-dong.import-on-startup=true --transit.admin-dong.file=<출력 GeoJSON>
"""

import argparse
import json
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory

import shapefile
from openpyxl import load_workbook
from pyproj import CRS, Transformer
from shapely.geometry import mapping, shape
from shapely.ops import transform

# 지도에 그리기 충분한 정밀도로 줄인다. 위도 37도 부근에서 0.00015도는 약 15m다.
DEFAULT_TOLERANCE_DEGREES = 0.00015


def code_part(value, width):
    if isinstance(value, float):
        value = int(value)
    return str(value).strip().zfill(width)


def read_region_names(path, sheet_name):
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook.worksheets[0]
    names = {}
    for sido_code, sido_name, sigungu_code, sigungu_name, dong_code, dong_name in sheet.iter_rows(
            min_row=3, max_col=6, values_only=True):
        if sido_code is None or dong_code is None:
            continue
        code = code_part(sido_code, 2) + code_part(sigungu_code, 3) + code_part(dong_code, 3)
        names[code] = {
            "sidoName": str(sido_name).strip(),
            "sigunguName": str(sigungu_name).strip(),
        }
    return names


def convert(shapefile_zip, region_codes, output, tolerance, sheet_name):
    names = read_region_names(region_codes, sheet_name)
    with TemporaryDirectory() as directory:
        with zipfile.ZipFile(shapefile_zip) as archive:
            archive.extractall(directory)
        base = Path(directory) / "BND_ADM_DONG_PG"
        source_crs = CRS.from_wkt(base.with_suffix(".prj").read_text())
        transformer = Transformer.from_crs(source_crs, "EPSG:4326", always_xy=True)
        reader = shapefile.Reader(str(base), encoding="cp949")
        fields = [field[0] for field in reader.fields[1:]]
        features = []
        for item in reader.iterShapeRecords():
            properties = dict(zip(fields, item.record))
            code = str(properties["ADM_CD"]).strip()
            name = properties["ADM_NM"].strip()
            region = names.get(code, {})
            geometry = transform(transformer.transform, shape(item.shape.__geo_interface__))
            geometry = geometry.simplify(tolerance, preserve_topology=True)
            features.append({
                "type": "Feature",
                "properties": {
                    "code": code,
                    "name": name,
                    "sidoName": region.get("sidoName", ""),
                    "sigunguName": region.get("sigunguName", ""),
                    "fullName": " ".join(filter(None, [region.get("sidoName"), region.get("sigunguName"), name])),
                    "baseDate": str(properties["BASE_DATE"]).strip(),
                },
                "geometry": mapping(geometry),
            })
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({"type": "FeatureCollection", "features": features},
                                 ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    missing = sum(1 for feature in features if not feature["properties"]["sidoName"])
    print(f"{len(features)}개 행정동을 {output}에 썼다. 지역 코드표에 없는 행정동: {missing}개")


def main():
    parser = argparse.ArgumentParser(description="SGIS 행정동 경계를 적재용 GeoJSON으로 바꾼다.")
    parser.add_argument("shapefile_zip", type=Path, help="SGIS 행정구역 경계(행정동) Shapefile zip")
    parser.add_argument("region_codes", type=Path, help="SGIS 센서스 공간정보 지역 코드 xlsx")
    parser.add_argument("output", type=Path, help="출력 GeoJSON 경로")
    parser.add_argument("--sheet", default=None, help="지역 코드 시트 이름. 생략하면 첫 시트")
    parser.add_argument("--tolerance", type=float, default=DEFAULT_TOLERANCE_DEGREES,
                        help="단순화 허용 오차(WGS84 도 단위)")
    arguments = parser.parse_args()
    convert(arguments.shapefile_zip, arguments.region_codes, arguments.output,
            arguments.tolerance, arguments.sheet)


if __name__ == "__main__":
    main()
